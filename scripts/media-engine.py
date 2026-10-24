"""Bounded local attachment decoding. Never executes attachment contents."""
import base64, io, json, sys, zipfile
from pathlib import Path
from xml.etree import ElementTree
sys.stdin.reconfigure(encoding='utf-8');sys.stdout.reconfigure(encoding='utf-8')
LIMIT=12000

def jpeg(image):
    image=image.convert('RGB');image.thumbnail((1280,1280))
    output=io.BytesIO();image.save(output,format='JPEG',quality=82)
    return {'type':'image','mimeType':'image/jpeg','data':base64.b64encode(output.getvalue()).decode()}

def unpacked_limit(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if sum(x.file_size for x in z.infolist())>30*1024*1024:raise ValueError('ARCHIVE_SIZE')

def run(req):
    raw=base64.b64decode(req['data'],validate=True)
    if len(raw)>30*1024*1024:raise ValueError('MEDIA_SIZE')
    kind=req['kind'];name=Path(req.get('name','')).name;ext=Path(name).suffix.lower()
    if kind=='image' or ext in ('.png','.jpg','.jpeg','.webp','.gif','.bmp'):
        from PIL import Image
        Image.MAX_IMAGE_PIXELS=20000000
        with Image.open(io.BytesIO(raw)) as img:
            if img.width*img.height>20000000:raise ValueError('IMAGE_SIZE')
            count=getattr(img,'n_frames',1);frames=[]
            for index in sorted(set([0,count//2,count-1] if count>1 else [0])):
                img.seek(index);frames.append(jpeg(img.copy()))
        return {'text':f'[图片或表情；动图抽样 {len(frames)} 帧]' if count>1 else '[图片或表情]','images':frames}
    if kind=='video' or ext in ('.mp4','.mov','.mkv','.webm','.avi'):
        import av
        frames=[];audio=[];audio_samples=0
        with av.open(io.BytesIO(raw)) as c:
            duration=float(c.duration or 0)/av.time_base
            if duration<=0 or duration>120:raise ValueError('VIDEO_DURATION')
            if not c.streams.video:raise ValueError('NO_VIDEO')
            for second in [duration*i/6 for i in range(6)]:
                c.seek(int(second*av.time_base),backward=True)
                for frame in c.decode(video=0):
                    if frame.time is None or frame.time>=second:
                        frames.append(jpeg(frame.to_image()));break
        # Speech is separately transcribed; frames are samples, not a full timeline.
        with av.open(io.BytesIO(raw)) as c:
            if c.streams.audio:
                resampler=av.AudioResampler(format='s16',layout='mono',rate=16000)
                for frame in c.decode(audio=0):
                    for sample in resampler.resample(frame):
                        b=sample.to_ndarray().tobytes();left=max(0,16000*2*60-audio_samples)
                        audio.append(b[:left]);audio_samples+=min(len(b),left)
                    if audio_samples>=16000*2*60:break
        result={'text':f'[视频 {duration:.1f} 秒：仅抽样 {len(frames)} 个画面；声音至多分析前60秒，不能推断未采样画面]','images':frames}
        if audio:
            import wave
            out=io.BytesIO()
            with wave.open(out,'wb') as w:w.setnchannels(1);w.setsampwidth(2);w.setframerate(16000);w.writeframes(b''.join(audio))
            result['audio']=base64.b64encode(out.getvalue()).decode()
        return result
    if ext in ('.txt','.md','.csv','.json','.log','.yaml','.yml'):
        if len(raw)>2*1024*1024:raise ValueError('TEXT_SIZE')
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:text=raw.decode('gb18030',errors='replace')
    elif ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:raise ValueError('PDF_ENCRYPTED')
        text=''
        for page in reader.pages[:20]:
            text+=(page.extract_text() or '')+'\n'
            if len(text)>=LIMIT:break
        if not text.strip():raise ValueError('PDF_SCANNED')
        text='[PDF：仅提取前20页内的文字，不含扫描页图片]\n'+text
    elif ext=='.docx':
        unpacked_limit(raw)
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            doc=ElementTree.fromstring(z.read('word/document.xml'))
            text='\n'.join(''.join(p.itertext()) for p in doc.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'))
    elif ext=='.xlsx':
        unpacked_limit(raw)
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        try:
            parts=[]
            for sheet in book.worksheets[:5]:
                parts.append('[工作表] '+sheet.title)
                for row in sheet.iter_rows(max_row=200,max_col=30,values_only=True):
                    parts.append('\t'.join(str(v) if v is not None else '' for v in row))
                    if sum(map(len,parts))>=LIMIT:break
            text='[表格抽样：最多5个工作表，各前200行、30列，公式只读取缓存值]\n'+'\n'.join(parts)
        finally:book.close()
    else:raise ValueError('FILE_TYPE')
    return {'text':f'[文件 {name[:200]}；提取文字，最多{LIMIT}字符]\n'+text[:LIMIT],'images':[]}

try:print(json.dumps(run(json.loads(sys.stdin.read(44000000))),ensure_ascii=False))
except Exception as exc:
    code=str(exc) if isinstance(exc,ValueError) and str(exc).isupper() else 'DECODE_FAILED'
    print(json.dumps({'error':code}));sys.exit(1)
